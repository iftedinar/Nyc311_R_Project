from pathlib import Path
import json,csv,re,hashlib,copy
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path.cwd()
C=json.loads(Path('analysis/worksheet_content.json').read_text())
def rows(name):
 with open(f'analysis/tables/{name}.csv',newline='',encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def fmt(v):
 try:
  x=float(v)
  return f'{x:,.0f}' if x.is_integer() else f'{x:,.4f}'
 except:return str(v)

source=Path('Project 1 Student Worksheet.docx')
D=Document(source); original_paragraphs=list(D.paragraphs)
from docx.enum.style import WD_STYLE_TYPE
for style_name,size in [('R Code',9),('R Output',8.5)]:
 s=D.styles.add_style(style_name,WD_STYLE_TYPE.PARAGRAPH)
 s.font.name='Consolas';s.font.size=Pt(size)
 s.paragraph_format.space_after=Pt(0)
 s.paragraph_format.line_spacing=1
Path('analysis/qa/artifact.md').write_text(f'''Reference: {source.resolve()}
SHA256: {hashlib.sha256(source.read_bytes()).hexdigest()}
Reference render: analysis/qa/reference.pdf (11 pages).
Page geometry: Letter portrait, all margins 0.7 inches. One section.
Normal: Calibri 10.5 pt. Question and part headings retained in original order.
Editable slots: all 31 numbered questions, response placeholders, 19 response tables,
name and checkbox paragraphs. Add question code/output after each answer group.
Original tables are retained and expanded as needed; 16-field missingness replaces
15 empty rows. Allow flow across pages with repeated headers. No fixed row heights.
User-authorized expansions: executed R examples and figures; optional LinkedIn appendix.
Remove non-question preliminary notes and unsourced citywide numbers; replace with
cited background in question 1. Normalize damaged replacement glyphs in prompt text.
Keep source file unchanged. Original styles, numbering and page geometry retained.
Code style: Consolas 9 pt, outputs 8.5 pt; paragraphs wrap rather than clip.
''')

# Replace only answer slots. Preserve all assignment questions and table locations.
for k,v in C['answers'].items():original_paragraphs[int(k)].text=v
for k in C['answers']:
 original_paragraphs[int(k)-1].paragraph_format.keep_with_next=True
original_paragraphs[2].text='Name: NYC 311 Analytics Student'
checks={122:'[ ] Minutes   [X] Hours   [ ] Days   [ ] Other',158:'[ ] No   [X] Yes - explain the transformation:',
188:'[ ] 0.50   [X] Other cutoff: 0.25',202:'[ ] Predicting slow when not slow   [X] Predicting not slow when actually slow',221:'[ ] No   [X] Yes'}
for i,t in checks.items():original_paragraphs[i].text=t
for i in range(8,20):
 p=original_paragraphs[i];p._element.getparent().remove(p._element)
# Add concise reading instructions before Part 1.
p=original_paragraphs[3].insert_paragraph_before('Run NYC311_Analysis.R from this folder before running the question examples. Each numbered question below includes its answers, an executable R example, and output from the completed R run. The full script reproduces every table, chart, model, and metric. Original source files are preserved. All results describe the supplied April 2025 sample.')
p=original_paragraphs[3].insert_paragraph_before('The R examples share objects created by the full script; run them in numerical order in the same R session. Full tables are in analysis/tables. Name is a permitted placeholder. Review the learning reflection in your own voice before submission.')

def fill_table(t,data):
 # data includes header; reuse original header style and widths.
 while len(t.rows)<len(data):t.add_row()
 while len(t.rows)>len(data):t._tbl.remove(t.rows[-1]._tr)
 for r,vals in zip(t.rows,data):
  for c,v in zip(r.cells,vals):c.text=str(v)

decisions={
'unique_key':('Remove predictor','Identifier; retain for audit'),
'created_date':('Derive features','Known at submission'),
'closed_date':('Target only','332 absent; exclude these rows'),
'agency':('Remove predictor','Largely determined by problem'),
'problem':('Keep and group','Pool rare training categories'),
'problem_detail':('Remove predictor','High cardinality and overlap'),
'additional_details':('Remove predictor','Incomplete free text'),
'status':('Remove predictor','Post-submission; all Closed'),
'location_type':('Remove predictor','Exact overlap with some problems'),
'incident_zip':('Remove predictor','Use borough for parsimony'),
'borough':('Keep','Five complete categories'),
'community_board':('Remove predictor','Geographic overlap; many levels'),
'open_data_channel_type':('Keep and recode','Rare level to training mode'),
'due_date':('Remove predictor','99.415% missing; timing uncertain'),
'resolution_description':('Remove predictor','Post-submission information'),
'resolution_action_updated_date':('Remove predictor','Post-submission information')}
tables=[]
tables.append([['Item','Answer'],['Number of rows before cleaning','40,000'],['Number of columns before cleaning','16'],['Number of rows after removing unusable records','39,668'],['Number of columns after selecting/creating variables for modeling','8 (6 predictors + 2 outcomes)']])
tables.append([['Variable','Number Missing','Percent Missing','Decision','Reason']]+[[r['variable'],fmt(r['missing']),f"{float(r['percent']):.3f}%",*decisions[r['variable']]] for r in rows('missingness')])
cat_table=[['Variable','Most common categories','Rare categories under 100','Modeling decision']]
for r in rows('category_summary'):
 v=r['variable'];freq=rows('frequency_'+v)[:3]
 cat_table.append([v,'; '.join(f"{a['category']} ({int(a['n']):,})" for a in freq),f"{r['rare_under_100']} of {r['levels']}",decisions[v][0]])
tables.append(cat_table)
tables.append([['Statistic','Value in hours unless count']]+[[r['statistic'],fmt(r['value'])] for r in rows('duration_summary')])
tables.append([['Variable','Comparison','Pattern','Model implication'],
 ['agency','Eligible group medians','EDC 7,254.84 h; NYPD 1.43 h','Large case-mix differences'],
 ['problem','Eligible types with n >= 100','Helicopter 7,254.84 h; Rodent 0 h','Retain problem; interpret closure cautiously'],
 ['borough','Median duration','Staten Island 13.07 h; Bronx 10.82 h','Context, not causal ranking'],
 ['channel','Median duration','UNKNOWN 45.02 h; PHONE 25.97 h','May reflect different issue mix']])
tables.append([['Cross-tab','Relationship','Finding'],['Agency by problem','Routing of problem types','138 of 142 types occur within one agency'],['Borough by channel','Within-borough channel shares','Phone: Bronx 34.66%, Manhattan 24.37%'],['Problem by borough','Request mix across geography','Full counts and row percentages exported']])
tables.append([['Key EDA finding','Why it matters'],['Strongly skewed duration','Median 6.68 h; mean 251.36 h. Use log1p and separate hour metrics.'],['332 missing closed dates','Cannot calculate an observed target; exclude without imputation.'],['Problem overlaps agency and location','Use a parsimonious predictor set; check training design rank.']])
tables.append([['Records removed','Reason'],['332 with missing closed_date','Cannot compute observed closure duration'],['0 duplicate keys','No repeated unique_key values'],['0 invalid created dates, negative durations, or outside-April rows','Checks implemented; no additional exclusions']])
tables.append([['New variable','How created','Purpose'],['created, closed','Parse local timestamp strings','Calculate elapsed time'],['resolution_hours','difftime(closed, created), hours','Raw target and EDA'],['log_resolution_hours','log1p(resolution_hours)','Linear regression outcome'],['day, hour, weekday','Extract creation components','EDA and time features'],['weekend','Saturday or Sunday indicator','EDA only; redundant with weekday'],['hour_sin, hour_cos','sin/cos(2*pi*hour/24)','Cyclic hourly predictors'],['slow','Hours exceed frozen training threshold','Binary logistic outcome'],['Grouped categorical values','Training counts and saved mapping','Stable sparse-level handling']])
tables.append([['Field','Decision','Reason'],['closed_date','Remove 332 rows','Do not impute the target'],['location_type','Explore as Missing; omit predictor','Exact training-design redundancy'],['problem_detail, additional_details, incident_zip, due_date','Omit these predictors','Sparsity, complexity, geography overlap, or timing'],['Final model predictors','No missing values after mappings','Verified in R']])
tables.append([['Categorical variable','Handling','Reason'],['problem','Training n < 100 pooled as Other','Limit sparse coefficients'],['borough','Retain all five levels','Coarse location context'],['open_data_channel_type','Keep common levels; rare/unseen to saved fallback','Single rare channel cannot support its own effect'],['weekday','Seven-level factor, Monday reference','Allow weekday differences']])
tables.append([['Variable','Why excluded from predictors'],['closed_date','Unknown when a request first arrives'],['status','Snapshot reflects later progress; all Closed here'],['resolution_description','Records agency action after submission'],['resolution_action_updated_date','Timestamp of later action'],['resolution_hours, log_resolution_hours, slow','Outcomes derived from future closure; never predictors'],['due_date','Availability at submission unverified; almost entirely missing']])
tables.append([['Predictor','Linear','Logistic','Reason'],['problem','Yes','Yes','Issue/workflow context'],['borough','Yes','Yes','Coarse geography'],['open_data_channel_type','Yes','Yes','Submission context'],['weekday','Yes','Yes','Day-of-week patterns'],['hour_sin','Yes','Yes','Cyclic time of submission'],['hour_cos','Yes','Yes','Completes cyclic representation']])
tables.append([['Item','Answer'],['Percent used for training','70% (27,767 / 39,668)'],['Percent used for validation','30% (11,901 / 39,668)'],['Training observations','27,767'],['Validation observations','11,901'],['Random seed','353']])
lm=rows('linear_metrics')
tables.append([['Metric on log1p-hours scale','Training','Validation']]+[[m,fmt(lm[0][m]),fmt(lm[1][m])] for m in ['RMSE','MAE','R_squared']])
tables.append([['','Predicted Not Slow','Predicted Slow'],['Actually Not Slow','4,997 (TN)','3,789 (FP)'],['Actually Slow','1,550 (FN)','1,565 (TP)']])
tables.append([['Metric','Validation value at 0.25'],['Accuracy','55.14%'],['Sensitivity / Recall','50.24%'],['Specificity','56.87%'],['Precision','29.23%'],['F1','0.3696'],['ROC AUC','0.5542'],['Flagged share','44.99%']])
tables.append([['Finding','Evidence'],['Duration varies widely','Median 6.68 h versus mean 251.36 h; problem/agency differences'],['Log fit does not ensure hour accuracy','Log RMSE 1.3089; hour RMSE 1,373.18 versus baseline 970.97'],['Slow classification is weak','AUC 0.5542; 50.24% recall, 29.23% precision; large borough imbalance']])
tables.append([['Limitation','Why it matters'],['One supplied month and unknown sampling method','Cannot generalize to other months or all city requests; boundary counts uneven'],['All records marked Closed','Still-open requests may be absent; 332 lack closure timestamps'],['Random within-month holdout','Does not test future periods or correlated related requests'],['Administrative closure outcome','Does not demonstrate issue resolution or resident satisfaction'],['Relative threshold and sparse-type fallback','Not an SLA; some rare types use a less comparable global cutoff'],['Weak classification and geographic imbalance','Automatic prioritization could allocate attention poorly']])
assert len(tables)==len(D.tables),(len(tables),len(D.tables))
for t,data in zip(D.tables,tables):fill_table(t,data)

# Insert code and executed output at the end of each numbered question.
qheads={int(m.group(1)):p for p in D.paragraphs if (m:=re.match(r'^(\d+)\. ',p.text)) and 1<=int(m.group(1))<=31}
figures={7:['time_day','time_hour'],8:['resolution_histograms'],9:['borough_duration'],21:['linear_diagnostics']}
def insert_before(anchor,text,style=None):return anchor.insert_paragraph_before(text,style)
for q,head in qheads.items():
 elems=list(D._element.body);start=elems.index(head._p); anchor=None
 for el in elems[start+1:]:
  if el.tag==qn('w:p'):
   txt=''.join(el.itertext())
   from docx.text.paragraph import Paragraph
   p=Paragraph(el,D)
   if re.match(r'^(\d+)\. |^Part \d|^Reminder:',p.text):anchor=p;break
 if anchor is None:anchor=D.add_paragraph()
 insert_before(anchor,f'R code for question {q}', 'Heading 3')
 for line in C['snippets'][str(q)].splitlines():
  p=insert_before(anchor,line);p.style='R Code'
 insert_before(anchor,'Executed R output','Heading 3')
 output=Path(f'analysis/question_output/Q{q:02}.txt').read_text()
 # Wide console data frames wrap safely; no generated or invented outputs.
 for line in output.rstrip().splitlines():
  p=insert_before(anchor,line);p.style='R Output'
 for fig in figures.get(q,[]):
  p=insert_before(anchor,'');p.add_run().add_picture(f'analysis/figures/{fig}.png',width=Inches(5.7 if q==9 else 6.8))

D.add_heading('Optional LinkedIn Article Analysis',level=1)
D.add_paragraph('Question: How does the probability cutoff change detection and staff workload, and does the same cutoff behave similarly across boroughs?')
D.add_paragraph('R code: run the Optional LinkedIn Article Analysis section of NYC311_Analysis.R. The following examples inspect its executed results.')
for line in ['read.csv("analysis/tables/optional_threshold_tradeoff.csv")','read.csv("analysis/tables/optional_borough_monitor.csv")']:
 D.add_paragraph(line,'R Code')
D.add_picture('analysis/figures/optional_threshold_tradeoff.png',width=Inches(6.8))
D.add_paragraph('At 0.25, the model flags approximately 450 of every 1,000 requests, with about 132 correct flags. It misses about half of actual slow requests. At 0.50 it flags none. A lower cutoff can increase recall while creating more review work; no optimal cutoff can be claimed without costs and capacity. These validation comparisons are exploratory, not an independent test of a tuned policy.')
D.add_paragraph('The same cutoff flags 99.43% of Bronx validation requests but only 1.17% in Staten Island. Recall is 99.35% and 1.09%, respectively. Different request mix and group calibration may contribute; these results do not establish discrimination, causality, or operational fairness. They show why a model can require subgroup review even when its overall metrics are available.')
D.add_heading('Sources and data definitions',level=1)
D.add_paragraph('Supplied sources: Project Overview- NYC 311 Service Requests.pdf; Project 1 Student Worksheet.docx; NYC311_April2025_RawSample.csv. No separate data dictionary or lecture script was included. The field-use table in question 5 documents the operational interpretation used for this analysis.')
D.add_paragraph('NYC 311 Reporting FAQ: https://www.nyc.gov/site/311reporting/faq/faq.page (accessed September 26, 2026). Defines the duration as time between receipt and agency-recorded closure.')
D.add_paragraph('NYC311 Service Requests: https://portal.311.nyc.gov/article/?kanumber=KA-03116 (accessed September 26, 2026). A closed service request may still have an unresolved underlying issue.')
D.add_paragraph('NYC Open Data dataset and column definitions: https://data.cityofnewyork.us/Social-Services/311-Service-Requests-from-2020-to-Present/erm2-nwe9 (accessed September 26, 2026). The supplied sample is the analysis input; the live dataset was not substituted.')

# Preserve source page system, improve expansion readability and remove broken glyphs.
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 D.styles[name].font.color.rgb=RGBColor(0,0,0)
for p in D.paragraphs:
 if '\ufffd' in p.text:
  p.text=p.text.replace('model\ufffds',"model's").replace('\ufffd','"')
 if p.style.name.startswith('Heading'):
  p.paragraph_format.keep_with_next=True
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
 if p.text.startswith('Part '):p.paragraph_format.page_break_before=False
 if p.style.name not in ['R Code','R Output'] and not p.style.name.startswith('Heading'):
  p.paragraph_format.space_after=Pt(5)
for t in D.tables:
 t.alignment=WD_TABLE_ALIGNMENT.CENTER
 t.autofit=False
 for i,r in enumerate(t.rows):
  trPr=r._tr.get_or_add_trPr()
  if i==0:trPr.append(OxmlElement('w:tblHeader'))
  for h in list(trPr.findall(qn('w:trHeight'))):trPr.remove(h)
  for c in r.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
    for run in p.runs:run.font.size=Pt(9);run.bold=(i==0)
for p in list(D.paragraphs):
 if not p.text and not p._p.findall('.//'+qn('w:drawing')):
  p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2)

out=Path('NYC311_Completed_Worksheet.docx');D.save(out)
assert '[Type your response here.]' not in '\n'.join(p.text for p in D.paragraphs)
Path('analysis/qa/worksheet_coverage.txt').write_text('\n'.join(f'Q{q:02}: question preserved, answer/table filled, R code and executed output included' for q in range(1,32)))
print(out.resolve())
